from __future__ import annotations

import hashlib
import os
import shutil
from collections import Counter
from collections.abc import Iterator
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from fraud_graph.config import AppConfig
from fraud_graph.data.manifest import DataManifest, QualityFinding, SourceFile
from fraud_graph.data.schemas import ACCOUNT_SCHEMA, LABEL_SCHEMA, TRANSACTION_SCHEMA
from fraud_graph.errors import DataContractError

LABEL_TEXT_TO_VALUE = {"其它": 0, "嫌疑人": 1, "受害人": 2}
INGESTION_VERSION = "xlsx-contract-v3"


def ingest_xlsx(source_dir: Path, config: AppConfig) -> DataManifest:
    source_paths = {
        "accounts": source_dir / config.source.account_workbook,
        "labels": source_dir / config.source.label_workbook,
        "transactions": source_dir / config.source.transaction_workbook,
    }
    for source_path in source_paths.values():
        if not source_path.is_file():
            raise DataContractError(f"required source file does not exist: {source_path}")

    source_files = [
        SourceFile(
            name=path.name,
            sha256=_sha256_file(path),
            size_bytes=path.stat().st_size,
        )
        for path in source_paths.values()
    ]
    version_material = "|".join(
        [
            f"ingestion:{INGESTION_VERSION}",
            f"schema:{config.schema_version}",
            *(f"{item.name}:{item.sha256}" for item in source_files),
        ]
    )
    data_version = hashlib.sha256(version_material.encode()).hexdigest()[:16]
    target_dir = config.storage.normalized_dir / data_version
    manifest_path = target_dir / "manifest.json"
    if manifest_path.exists():
        return DataManifest.load(manifest_path)

    temporary_dir = config.storage.normalized_dir / f".tmp-{data_version}-{os.getpid()}"
    if temporary_dir.exists():
        _remove_owned_temporary_dir(temporary_dir, config.storage.normalized_dir)
    temporary_dir.mkdir(parents=True, exist_ok=False)

    try:
        accounts, account_labels = _read_accounts(source_paths["accounts"])
        labels, label_counts, label_findings = _read_labels(
            source_paths["labels"], account_labels, ZoneInfo(config.source.timezone)
        )
        pq.write_table(
            pa.Table.from_pylist(accounts, schema=ACCOUNT_SCHEMA),
            temporary_dir / "accounts.parquet",
            compression="zstd",
        )
        pq.write_table(
            pa.Table.from_pylist(labels, schema=LABEL_SCHEMA),
            temporary_dir / "labels.parquet",
            compression="zstd",
        )

        transaction_summary = _write_transactions(
            source_paths["transactions"],
            temporary_dir / "transactions",
            set(account_labels),
            source_files[2].sha256[:12],
            ZoneInfo(config.source.timezone),
            config.ingestion.batch_size,
        )
        findings = [*label_findings, *transaction_summary["findings"]]
        active_by_label = Counter(
            account_labels[account_id]
            for account_id in transaction_summary["active_accounts"]
            if account_id in account_labels
        )
        for label, total in Counter(account_labels.values()).items():
            active = active_by_label[label]
            if label != 0 and active < total:
                findings.append(
                    QualityFinding(
                        code="RISK_LABEL_WITHOUT_TRANSACTION",
                        severity="WARN",
                        message=f"label {label} has {total - active} accounts without transactions",
                        observed=total - active,
                    )
                )
        temporal_training_ready = _temporal_labels_ready(labels)
        if not temporal_training_ready:
            findings.append(
                QualityFinding(
                    code="LABEL_TIME_MISSING",
                    severity="BLOCKER",
                    message=(
                        "risk_event_at and label_available_at are required for temporal training; "
                        "the current dataset supports retrospective analysis only"
                    ),
                )
            )

        manifest = DataManifest(
            schema_version=config.schema_version,
            ingestion_version=INGESTION_VERSION,
            data_version=data_version,
            created_at=datetime.now(UTC),
            source_timezone=config.source.timezone,
            sources=source_files,
            account_rows=len(accounts),
            label_rows=len(labels),
            transaction_rows=transaction_summary["rows"],
            transaction_time_min=transaction_summary["time_min"],
            transaction_time_max=transaction_summary["time_max"],
            unique_active_accounts=len(transaction_summary["active_accounts"]),
            label_counts={str(key): value for key, value in sorted(label_counts.items())},
            temporal_training_ready=temporal_training_ready,
            findings=findings,
            files={
                "accounts": "accounts.parquet",
                "labels": "labels.parquet",
                "transactions": "transactions/year=*/month=*/*.parquet",
            },
        )
        (temporary_dir / "manifest.json").write_text(
            manifest.model_dump_json(indent=2), encoding="utf-8"
        )
        target_dir.parent.mkdir(parents=True, exist_ok=True)
        temporary_dir.rename(target_dir)
        return manifest
    except Exception:
        _remove_owned_temporary_dir(temporary_dir, config.storage.normalized_dir)
        raise


def _read_accounts(path: Path) -> tuple[list[dict[str, Any]], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = tuple(next(rows))
        expected = ("账户脱敏id", "是否风险账户标签", "开户时长", "地区编码", "客户类型")
        if header != expected:
            raise DataContractError(f"unexpected account header: {header!r}")
        records: list[dict[str, Any]] = []
        labels: dict[str, int] = {}
        for source_row, row in enumerate(rows, start=2):
            account_raw, label_raw, tenure_raw, region_raw, customer_type_raw = row
            if any(value is None for value in row):
                raise DataContractError(f"null value in account row {source_row}")
            account_id = _opaque_id(account_raw)
            if account_id in labels:
                raise DataContractError(f"duplicate account_id {account_id} at row {source_row}")
            label = int(label_raw)
            if label not in (0, 1, 2):
                raise DataContractError(f"invalid account label {label} at row {source_row}")
            labels[account_id] = label
            records.append(
                {
                    "account_id": account_id,
                    "tenure_months": float(tenure_raw),
                    "region_code": _opaque_id(region_raw),
                    "customer_type": str(customer_type_raw).strip(),
                    "source_row": source_row,
                }
            )
        return records, labels
    finally:
        workbook.close()


def _read_labels(
    path: Path,
    account_labels: dict[str, int],
    source_timezone: ZoneInfo,
) -> tuple[list[dict[str, Any]], Counter[int], list[QualityFinding]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = tuple(next(rows))
        if len(header) < 2 or header[0] != "账户脱敏id" or header[1] != "标签类型":
            raise DataContractError(f"unexpected label header: {header!r}")
        optional_columns = {
            "risk_event_at": _first_header_index(
                header, "风险事件时间", "风险事件时间分桶", "risk_event_at"
            ),
            "label_available_at": _first_header_index(
                header, "标签生效时间", "标签时间分桶", "label_available_at"
            ),
            "label_source": _first_header_index(header, "标签来源类别", "label_source"),
            "case_id": _first_header_index(header, "案件脱敏ID", "案件ID", "case_id"),
        }
        records: list[dict[str, Any]] = []
        seen: set[str] = set()
        counts: Counter[int] = Counter()
        findings: list[QualityFinding] = []
        for source_row, row in enumerate(rows, start=2):
            account_raw, label_text_raw = row[:2]
            if account_raw is None or label_text_raw is None:
                raise DataContractError(f"null value in label row {source_row}")
            account_id = _opaque_id(account_raw)
            label_text = str(label_text_raw).strip()
            if account_id in seen:
                raise DataContractError(f"duplicate label for account {account_id}")
            if label_text not in LABEL_TEXT_TO_VALUE:
                raise DataContractError(f"unsupported label type {label_text!r}")
            if account_id not in account_labels:
                raise DataContractError(f"label references unknown account {account_id}")
            label = LABEL_TEXT_TO_VALUE[label_text]
            if account_labels[account_id] != label:
                raise DataContractError(f"account and label workbook disagree for {account_id}")
            seen.add(account_id)
            counts[label] += 1
            records.append(
                {
                    "account_id": account_id,
                    "label": label,
                    "label_type": label_text,
                    "risk_event_at": _optional_datetime(
                        _row_value(row, optional_columns["risk_event_at"]),
                        source_timezone,
                        "risk_event_at",
                        source_row,
                    ),
                    "label_available_at": _optional_datetime(
                        _row_value(row, optional_columns["label_available_at"]),
                        source_timezone,
                        "label_available_at",
                        source_row,
                    ),
                    "label_source": _optional_text(
                        _row_value(row, optional_columns["label_source"])
                    ),
                    "case_id": _optional_text(_row_value(row, optional_columns["case_id"])),
                    "source_row": source_row,
                }
            )
        missing = set(account_labels) - seen
        if missing:
            findings.append(
                QualityFinding(
                    code="ACCOUNT_WITHOUT_LABEL",
                    severity="ERROR",
                    message="accounts are missing risk labels",
                    observed=len(missing),
                )
            )
        return records, counts, findings
    finally:
        workbook.close()


def _temporal_labels_ready(labels: list[dict[str, Any]]) -> bool:
    return all(
        row["label_available_at"] is not None
        and (row["label"] == 0 or row["risk_event_at"] is not None)
        for row in labels
    )


def _first_header_index(header: tuple[Any, ...], *names: str) -> int | None:
    for name in names:
        if name in header:
            return header.index(name)
    return None


def _row_value(row: tuple[Any, ...], index: int | None) -> Any:
    return None if index is None or index >= len(row) else row[index]


def _optional_text(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    return str(value).strip()


def _optional_bool(value: Any) -> bool | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "是", "y", "yes"}:
        return True
    if normalized in {"0", "false", "否", "n", "no"}:
        return False
    raise DataContractError(f"invalid boolean value: {value!r}")


def _parse_transaction_datetime(value: Any, source_row: int) -> datetime:
    if isinstance(value, datetime):
        return value
    raw = str(value).strip()
    for pattern in (
        "%Y%m%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y%m%d",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(raw, pattern)
        except ValueError:
            continue
    raise DataContractError(f"invalid transaction time in row {source_row}: {value!r}")


def _optional_datetime(
    value: Any,
    source_timezone: ZoneInfo,
    field_name: str,
    source_row: int,
) -> datetime | None:
    if value is None or str(value).strip() == "":
        return None
    parsed_datetime: datetime | None = None
    if isinstance(value, datetime):
        parsed_datetime = value
    else:
        raw = str(value).strip()
        parsed: datetime | None = None
        for pattern in (
            "%Y%m%d %H:%M:%S",
            "%Y-%m-%d %H:%M:%S",
            "%Y%m%d",
            "%Y-%m-%d",
        ):
            try:
                parsed = datetime.strptime(raw, pattern)
                break
            except ValueError:
                continue
        if parsed is None:
            raise DataContractError(f"invalid {field_name} in label row {source_row}: {value!r}")
        parsed_datetime = parsed
    if parsed_datetime is None:
        raise DataContractError(f"invalid {field_name} in label row {source_row}")
    if parsed_datetime.tzinfo is None:
        return parsed_datetime.replace(tzinfo=source_timezone).astimezone(UTC)
    return parsed_datetime.astimezone(UTC)


def _write_transactions(
    source_path: Path,
    output_dir: Path,
    account_ids: set[str],
    source_hash_prefix: str,
    source_timezone: ZoneInfo,
    batch_size: int,
) -> dict[str, Any]:
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    writers: dict[tuple[int, int], pq.ParquetWriter] = {}
    buffers: dict[tuple[int, int], list[dict[str, Any]]] = {}
    findings: list[QualityFinding] = []
    active_accounts: set[str] = set()
    rows_written = 0
    self_transfers = 0
    negative_amounts = 0
    time_min: datetime | None = None
    time_max: datetime | None = None
    try:
        rows: Iterator[tuple[Any, ...]] = workbook.active.iter_rows(values_only=True)
        header = tuple(next(rows))
        required = ("付款账户脱敏id", "收款账户脱敏id", "交易时间", "金额")
        if tuple(header[: len(required)]) != required:
            raise DataContractError(f"unexpected transaction header: {header!r}")
        optional_columns = {
            "channel_type": _first_header_index(header, "渠道类型", "channel_type"),
            "transaction_type": _first_header_index(header, "交易类型", "transaction_type"),
            "debit_credit_direction": _first_header_index(
                header, "借贷方向", "收支方向", "debit_credit_direction"
            ),
            "is_cross_region": _first_header_index(
                header, "是否跨区", "是否跨区域", "is_cross_region"
            ),
        }
        for source_row, row in enumerate(rows, start=2):
            required_values = row[: len(required)]
            if len(required_values) < len(required) or any(
                value is None for value in required_values
            ):
                raise DataContractError(f"null value in transaction row {source_row}")
            payer_raw, payee_raw, event_time_raw, amount_raw = required_values
            payer = _opaque_id(payer_raw)
            payee = _opaque_id(payee_raw)
            if payer not in account_ids or payee not in account_ids:
                raise DataContractError(
                    f"transaction row {source_row} references an unknown account"
                )
            event_time_local = _parse_transaction_datetime(event_time_raw, source_row)
            event_time_utc = event_time_local.replace(tzinfo=source_timezone).astimezone(UTC)
            try:
                amount = Decimal(str(amount_raw)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            except InvalidOperation as exc:
                raise DataContractError(f"invalid amount in row {source_row}") from exc
            if abs(amount) >= Decimal("1000000000000000000"):
                raise DataContractError(f"amount exceeds decimal contract in row {source_row}")

            partition = (event_time_local.year, event_time_local.month)
            buffers.setdefault(partition, []).append(
                {
                    "transaction_id": f"tx_{source_hash_prefix}_{source_row:09d}",
                    "payer_account_id": payer,
                    "payee_account_id": payee,
                    "event_time_utc": event_time_utc,
                    "event_time_local": event_time_local,
                    "amount": amount,
                    "is_negative": amount < 0,
                    "is_self_transfer": payer == payee,
                    "channel_type": _optional_text(
                        _row_value(row, optional_columns["channel_type"])
                    ),
                    "transaction_type": _optional_text(
                        _row_value(row, optional_columns["transaction_type"])
                    ),
                    "debit_credit_direction": _optional_text(
                        _row_value(row, optional_columns["debit_credit_direction"])
                    ),
                    "is_cross_region": _optional_bool(
                        _row_value(row, optional_columns["is_cross_region"])
                    ),
                    "source_row": source_row,
                }
            )
            active_accounts.update((payer, payee))
            self_transfers += int(payer == payee)
            negative_amounts += int(amount < 0)
            time_min = event_time_utc if time_min is None or event_time_utc < time_min else time_min
            time_max = event_time_utc if time_max is None or event_time_utc > time_max else time_max
            rows_written += 1
            if len(buffers[partition]) >= batch_size:
                _flush_partition(partition, buffers, writers, output_dir)
        for partition in list(buffers):
            _flush_partition(partition, buffers, writers, output_dir)
    finally:
        workbook.close()
        for writer in writers.values():
            writer.close()

    if self_transfers:
        findings.append(
            QualityFinding(
                code="SELF_TRANSFER_PRESENT",
                severity="WARN",
                message="self transfers require an explicit business interpretation",
                observed=self_transfers,
            )
        )
    if negative_amounts:
        findings.append(
            QualityFinding(
                code="NEGATIVE_AMOUNT_PRESENT",
                severity="WARN",
                message="negative amounts require reversal/refund semantics",
                observed=negative_amounts,
            )
        )
    return {
        "rows": rows_written,
        "active_accounts": active_accounts,
        "time_min": time_min,
        "time_max": time_max,
        "findings": findings,
    }


def _flush_partition(
    partition: tuple[int, int],
    buffers: dict[tuple[int, int], list[dict[str, Any]]],
    writers: dict[tuple[int, int], pq.ParquetWriter],
    output_dir: Path,
) -> None:
    records = buffers[partition]
    if not records:
        return
    year, month = partition
    partition_dir = output_dir / f"year={year:04d}" / f"month={month:02d}"
    partition_dir.mkdir(parents=True, exist_ok=True)
    writer = writers.get(partition)
    if writer is None:
        writer = pq.ParquetWriter(
            partition_dir / "part-00000.parquet",
            TRANSACTION_SCHEMA,
            compression="zstd",
        )
        writers[partition] = writer
    writer.write_table(pa.Table.from_pylist(records, schema=TRANSACTION_SCHEMA))
    buffers[partition] = []


def _opaque_id(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while block := source.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def _remove_owned_temporary_dir(path: Path, parent: Path) -> None:
    resolved_path = path.resolve()
    resolved_parent = parent.resolve()
    if resolved_path.parent != resolved_parent or not resolved_path.name.startswith(".tmp-"):
        raise RuntimeError(f"refusing to remove non-owned path: {path}")
    shutil.rmtree(resolved_path)
