from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from fraud_graph.config import AppConfig


@pytest.fixture()
def source_dir(tmp_path: Path) -> Path:
    source = tmp_path / "source"
    source.mkdir()

    accounts = Workbook()
    sheet = accounts.active
    sheet.append(["账户脱敏id", "是否风险账户标签", "开户时长", "地区编码", "客户类型"])
    for account_id in range(1, 21):
        label = 1 if account_id in {1, 2, 3, 4} else 0
        sheet.append([account_id, str(label), float(account_id + 5), account_id % 3, "对私"])
    accounts.save(source / "账户表.xlsx")

    labels = Workbook()
    sheet = labels.active
    sheet.append(["账户脱敏id", "标签类型"])
    for account_id in range(1, 21):
        sheet.append([account_id, "嫌疑人" if account_id in {1, 2, 3, 4} else "其它"])
    labels.save(source / "风险标签表.xlsx")

    transactions = Workbook(write_only=True)
    sheet = transactions.create_sheet("Sheet0")
    sheet.append(["付款账户脱敏id", "收款账户脱敏id", "交易时间", "金额"])
    rows = [
        [1, 5, "20250101 10:00:00", 1000.0],
        [5, 6, "20250101 10:10:00", 900.0],
        [6, 7, "20250101 10:20:00", 850.0],
        [8, 2, "20250101 11:00:00", 2000.0],
        [9, 8, "20250101 10:40:00", 1800.0],
        [10, 10, "20250102 12:00:00", -50.0],
    ]
    for row in rows:
        sheet.append(row)
    transactions.save(source / "交易边表.xlsx")
    return source


@pytest.fixture()
def config(tmp_path: Path) -> AppConfig:
    return AppConfig.model_validate(
        {
            "schema_version": "1.0",
            "source": {
                "timezone": "Asia/Shanghai",
                "account_workbook": "账户表.xlsx",
                "label_workbook": "风险标签表.xlsx",
                "transaction_workbook": "交易边表.xlsx",
            },
            "storage": {
                "normalized_dir": tmp_path / "normalized",
                "feature_dir": tmp_path / "features",
                "artifact_dir": tmp_path / "artifacts",
            },
            "ingestion": {"batch_size": 1000},
            "features": {"lookback_days": 30},
            "model": {
                "positive_labels": [1],
                "random_seed": 42,
                "test_size": 0.2,
                "top_fraction": 0.1,
            },
            "paths": {},
        }
    )


@pytest.fixture()
def temporal_source_dir(tmp_path: Path) -> Path:
    source = tmp_path / "temporal-source"
    source.mkdir()

    accounts = Workbook()
    account_sheet = accounts.active
    account_sheet.append(["账户脱敏id", "是否风险账户标签", "开户时长", "地区编码", "客户类型"])
    for account_id in range(1, 7):
        account_sheet.append(
            [account_id, "1" if account_id <= 4 else "0", 12.0, account_id % 2, "对私"]
        )
    accounts.save(source / "账户表.xlsx")

    labels = Workbook()
    label_sheet = labels.active
    label_sheet.append(
        ["账户脱敏id", "标签类型", "风险事件时间", "标签生效时间", "标签来源类别", "案件脱敏ID"]
    )
    risk_events = {1: "20250106", 2: "20250110", 3: "20250114", 4: "20250118"}
    for account_id in range(1, 7):
        if account_id <= 4:
            event = risk_events[account_id]
            available = (datetime.strptime(event, "%Y%m%d") + timedelta(days=1)).strftime("%Y%m%d")
            label_sheet.append(
                [account_id, "嫌疑人", event, available, "CASE", f"case-{account_id}"]
            )
        else:
            label_sheet.append([account_id, "其它", None, "20250101", "SYSTEM", None])
    labels.save(source / "风险标签表.xlsx")

    transactions = Workbook()
    transaction_sheet = transactions.active
    transaction_sheet.append(["付款账户脱敏id", "收款账户脱敏id", "交易时间", "金额"])
    for day in range(1, 21):
        transaction_sheet.append([5, 6, f"202501{day:02d} 10:00:00", 100 + day])
        transaction_sheet.append([1 + ((day - 1) % 4), 5, f"202501{day:02d} 11:00:00", 50 + day])
    transactions.save(source / "交易边表.xlsx")
    return source


@pytest.fixture()
def temporal_config(tmp_path: Path) -> AppConfig:
    return AppConfig.model_validate(
        {
            "schema_version": "1.0",
            "source": {
                "timezone": "Asia/Shanghai",
                "account_workbook": "账户表.xlsx",
                "label_workbook": "风险标签表.xlsx",
                "transaction_workbook": "交易边表.xlsx",
            },
            "storage": {
                "normalized_dir": tmp_path / "temporal-normalized",
                "feature_dir": tmp_path / "temporal-features",
                "artifact_dir": tmp_path / "temporal-artifacts",
            },
            "ingestion": {"batch_size": 1000},
            "temporal": {
                "observation_days": 3,
                "prediction_days": 2,
                "label_maturity_days": 1,
                "anchor_step_days": 2,
                "train_fraction": 0.6,
                "validation_fraction": 0.2,
            },
            "model": {
                "positive_labels": [1],
                "random_seed": 42,
                "top_fraction": 0.5,
            },
            "paths": {},
        }
    )
