from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import pyarrow.parquet as pq
from openpyxl import load_workbook

from fraud_graph.config import AppConfig
from fraud_graph.data.ingest import ingest_xlsx
from fraud_graph.features.builder import build_account_features


def test_ingest_is_versioned_and_idempotent(source_dir: Path, config: AppConfig) -> None:
    first = ingest_xlsx(source_dir, config)
    second = ingest_xlsx(source_dir, config)

    assert first.data_version == second.data_version
    assert first.account_rows == 20
    assert first.transaction_rows == 6
    assert first.temporal_training_ready is False
    assert {item.code for item in first.findings} >= {
        "LABEL_TIME_MISSING",
        "NEGATIVE_AMOUNT_PRESENT",
        "SELF_TRANSFER_PRESENT",
    }
    data_dir = config.storage.normalized_dir / first.data_version
    assert (data_dir / "manifest.json").is_file()
    assert pq.read_table(data_dir / "accounts.parquet").num_rows == 20


def test_features_respect_as_of(source_dir: Path, config: AppConfig) -> None:
    data = ingest_xlsx(source_dir, config)
    features = build_account_features(
        config.storage.normalized_dir,
        config.storage.feature_dir,
        data.data_version,
        lookback_days=30,
    )
    rows = pq.read_table(
        config.storage.feature_dir / features.feature_version / features.file
    ).to_pylist()
    by_account = {row["account_id"]: row for row in rows}

    assert features.row_count == 20
    assert by_account["1"]["out_txn_count"] == 1
    assert by_account["5"]["same_day_in_out_days"] == 1
    assert by_account["20"]["has_transactions"] == 0


def test_future_transactions_cannot_change_as_of_features(
    source_dir: Path, config: AppConfig
) -> None:
    as_of = datetime(2025, 1, 3, tzinfo=UTC)
    original_data = ingest_xlsx(source_dir, config)
    original_features = build_account_features(
        config.storage.normalized_dir,
        config.storage.feature_dir,
        original_data.data_version,
        as_of=as_of,
        lookback_days=30,
    )
    original_rows = pq.read_table(
        config.storage.feature_dir / original_features.feature_version / original_features.file
    ).to_pylist()

    transaction_path = source_dir / "交易边表.xlsx"
    workbook = load_workbook(transaction_path)
    sheet = workbook.active
    sheet.append([1, 20, "20250103 08:00:00", 999_999.0])
    sheet.append([1, 20, "20250201 08:00:00", 999_999.0])
    workbook.save(transaction_path)
    workbook.close()

    perturbed_data = ingest_xlsx(source_dir, config)
    perturbed_features = build_account_features(
        config.storage.normalized_dir,
        config.storage.feature_dir,
        perturbed_data.data_version,
        as_of=as_of,
        lookback_days=30,
    )
    perturbed_rows = pq.read_table(
        config.storage.feature_dir / perturbed_features.feature_version / perturbed_features.file
    ).to_pylist()

    assert perturbed_data.data_version != original_data.data_version
    assert perturbed_rows == original_rows


def test_optional_transaction_columns_are_parsed(source_dir: Path, config: AppConfig) -> None:
    transaction_path = source_dir / "交易边表.xlsx"
    workbook = load_workbook(transaction_path)
    sheet = workbook.active
    sheet.cell(row=1, column=5, value="渠道类型")
    sheet.cell(row=1, column=6, value="交易类型")
    sheet.cell(row=1, column=7, value="借贷方向")
    sheet.cell(row=1, column=8, value="是否跨区")
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=5, value="网银")
        sheet.cell(row=row, column=6, value="转账")
        sheet.cell(row=row, column=7, value="出")
        sheet.cell(row=row, column=8, value="是")
    workbook.save(transaction_path)
    workbook.close()

    data = ingest_xlsx(source_dir, config)
    transaction_file = next(
        (config.storage.normalized_dir / data.data_version / "transactions").glob(
            "year=*/month=*/*.parquet"
        )
    )
    row = pq.read_table(transaction_file).to_pylist()[0]
    assert row["channel_type"] == "网银"
    assert row["transaction_type"] == "转账"
    assert row["debit_credit_direction"] == "出"
    assert row["is_cross_region"] is True
